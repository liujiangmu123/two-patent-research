# -*- coding: utf-8 -*-
"""临时脚本：读取用户在先申请权利要求书（python-docx），输出纯文本到 _tmp。"""
import sys, pathlib
import docx

src = pathlib.Path(r"h:\Axinjihua\02动画项目\05动画Harness工作台\专利文档资料\专利\InSAR-GNSS联合监测装置\08_发明专利申请_V12_立杆加粗与中值平衡\申请文件\权利要求书.docx")
out = pathlib.Path(__file__).with_name("prior_claims.txt")
d = docx.Document(str(src))
lines = [p.text for p in d.paragraphs if p.text.strip()]
for t in d.tables:
    for r in t.rows:
        lines.append(" | ".join(c.text for c in r.cells))
out.write_text("\n".join(lines), encoding="utf-8")
print("paragraphs:", len(lines), "chars:", sum(len(x) for x in lines))
