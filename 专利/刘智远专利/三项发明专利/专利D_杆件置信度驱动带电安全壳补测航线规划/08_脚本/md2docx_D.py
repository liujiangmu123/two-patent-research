import re,sys,json,os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
B=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
J=os.path.join(B,"04_仿真验证","关键结果摘要.json")
vals=json.load(open(J,encoding="utf8")) if os.path.exists(J) else {}
def fill(s): return re.sub(r"【待填:([^】]+)】",lambda m:str(vals[m.group(1)]) if m.group(1) in vals else m.group(0),s)
def conv(md):
    d=Document(); st=d.styles["Normal"]; st.font.name="宋体"; st.font.size=Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
    lines=fill(open(md,encoding="utf8").read()).splitlines(); i=0
    while i<len(lines):
        l=lines[i]
        if l.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                c=[x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",x) for x in c if x): rows.append(c)
                i+=1
            n=max(map(len,rows)); t=d.add_table(rows=len(rows),cols=n); t.style="Table Grid"
            for r,row in enumerate(rows):
                for c,x in enumerate(row): t.cell(r,c).text=x.replace("`","")
            continue
        m=re.match(r"(#+) (.*)",l)
        if m: d.add_heading(m.group(2),min(len(m.group(1))-1,4) if len(m.group(1))>1 else 0)
        elif l.strip(): d.add_paragraph(re.sub(r"\*\*|`","",l.lstrip("> ")))
        i+=1
    d.save(md[:-3]+".docx"); print("ok",md[:-3]+".docx")
for p in ["01_调研与查新/查新与创新点分析报告.md","01_调研与查新/检索记录.md","02_技术交底书/技术交底书.md","07_申请文件/权利要求书.md","07_申请文件/说明书.md","07_申请文件/说明书摘要.md"]:
    conv(os.path.join(B,p))
