# md -> docx（标题/段落/表格/粗体），用法: python md2docx.py a.md [b.md ...]
import re,sys,json,os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
def fill(text):
    p=os.path.join(os.path.dirname(__file__),"..","04_仿真验证","关键结果摘要.json")
    if not os.path.exists(p): return text
    d=json.load(open(p,encoding="utf-8"))
    flat={}
    def walk(o,pre=""):
        if isinstance(o,dict):
            for k,v in o.items(): walk(v,k)
        else: flat[pre]=o
    walk(d)
    return re.sub(r"【待填:([^】]+)】",lambda m:str(flat.get(m.group(1),m.group(0))),text)
def run(p,t):
    for i,seg in enumerate(re.split(r"\*\*",t)):
        r=p.add_run(seg.replace("`","")); r.bold=bool(i%2)
def conv(md):
    doc=Document(); st=doc.styles["Normal"]; st.font.name="宋体"; st.font.size=Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
    lines=fill(open(md,encoding="utf-8").read()).split("\n"); i=0
    while i<len(lines):
        l=lines[i]
        if l.startswith("<!--") or not l.strip(): i+=1; continue
        m=re.match(r"(#+)\s+(.*)",l)
        if m: doc.add_heading(m.group(2),min(len(m.group(1))-1,4) or 0); i+=1; continue
        if l.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                c=[x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",x) for x in c): rows.append(c)
                i+=1
            n=max(len(r) for r in rows); t=doc.add_table(rows=len(rows),cols=n); t.style="Table Grid"
            for a,r in enumerate(rows):
                for b,x in enumerate(r): t.cell(a,b).text=""; run(t.cell(a,b).paragraphs[0],x)
            continue
        if l.startswith("> "): l=l[2:]
        if re.match(r"\s*[-*] ",l): run(doc.add_paragraph(style="List Bullet"),re.sub(r"^\s*[-*] ","",l))
        else: run(doc.add_paragraph(),l.strip())
        i+=1
    out=md[:-3]+".docx"; doc.save(out); print(out)
for f in sys.argv[1:]: conv(f)
