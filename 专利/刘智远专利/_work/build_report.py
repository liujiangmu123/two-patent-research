# -*- coding: utf-8 -*-
"""Assemble the Chinese analysis & patent-mining report (Word)."""
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import report_content_1 as c1  # noqa: E402
import report_content_2 as c2  # noqa: E402
from report_docx_lib import ACCENT, add_inline, add_table, new_doc, render  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")
OUT = Path(__file__).resolve().parent.parent / "PTM论文深度分析与专利挖掘研究报告.docx"


def cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "输电塔 LoD3 桁架重建（PTM）论文", size=22, color=ACCENT, base_bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "深度分析与专利挖掘研究报告", size=22, color=ACCENT, base_bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    add_inline(p, "—— 基于 IEEE TII 2026《A Novel Power Pylon Truss Modeler for LoD3 Reconstruction From ULS Data》",
               size=11.5, color=RGBColor(0x40, 0x40, 0x40))
    for _ in range(6):
        doc.add_paragraph()
    add_table(doc, ["项目", "内容"], [
        ["面向对象", "刘智远专利组（兰州石化职业技术大学）"],
        ["研究范围", "论文精读与复核；2024–2026 国内外研究进展；中国专利现有技术扫描；专利候选与权利要求草案；申报策略"],
        ["编制日期", "2026-09-30"],
        ["版本", "V1.0（内部研究资料，权利要求为草案，需专利代理师定稿）"],
    ], [3.2, 12.0], 10)
    doc.add_page_break()


def main():
    doc = new_doc()
    cover(doc)
    p = doc.add_paragraph()
    add_inline(p, "目  录", size=15, color=ACCENT, base_bold=True)
    render(doc, [("toc",)])
    p = doc.add_paragraph()
    add_inline(p, "（如目录未显示页码，请在 Word 中右键目录选择“更新域”。）", size=8.5, color=RGBColor(0x80, 0x80, 0x80))
    doc.add_page_break()
    render(doc, c1.SUMMARY)
    doc.add_page_break()
    render(doc, c1.PAPER_INFO + c1.PROBLEM + c1.METHOD + c1.RESULTS + c1.CRITIQUE)
    render(doc, c1.FRONTIER + c1.CN_SCAN)
    render(doc, c2.blocks())
    doc.save(OUT)
    print("saved:", OUT.name, round(OUT.stat().st_size / 1e6, 2), "MB")


if __name__ == "__main__":
    main()
