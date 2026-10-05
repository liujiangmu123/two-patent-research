# -*- coding: utf-8 -*-
"""Small python-docx helper layer: block list -> formatted Chinese report."""
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT, HEAD_FONT, LATIN = "宋体", "黑体", "Times New Roman"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def _set_font(run, east=BODY_FONT, latin=LATIN, size=None, bold=None, color=None, italic=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_inline(p, text, size=None, color=None, base_bold=False):
    """Supports **bold** and `mono` inline markup."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2])
            _set_font(r, size=size, bold=True, color=color)
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            _set_font(r, east=BODY_FONT, latin="Consolas", size=(size or 10.5) - 1, color=color)
        else:
            r = p.add_run(part)
            _set_font(r, size=size, bold=base_bold or None, color=color)


def add_hyperlink(p, text, url, size=9):
    part = p.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "1F5FAF"), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rpr.append(el)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), LATIN)
    fonts.set(qn("w:hAnsi"), LATIN)
    rpr.append(fonts)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _borders(table, color="8EA9C8"):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _cell_margins(table, top=50, bottom=50, left=80, right=80):
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_table(doc, headers, rows, widths_cm, size=8.5, head_fill="DCE6F1", zebra="F7F9FC"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    _cell_margins(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Cm(widths_cm[i])
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, h, size=size, base_bold=True)
        _shade(c, head_fill)
    _repeat_header(t.rows[0])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            c = cells[i]
            c.width = Cm(widths_cm[i])
            c.text = ""
            lines = str(v).split("\n")
            for li, line in enumerate(lines):
                p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                add_inline(p, line, size=size)
            if zebra and ri % 2 == 1:
                _shade(c, zebra)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    st = doc.styles["Normal"]
    st.font.name = LATIN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.3
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        hs = doc.styles[name]
        hs.font.name = LATIN
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = ACCENT
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        hs.paragraph_format.space_before = Pt(10 if name != "Heading 3" else 6)
        hs.paragraph_format.space_after = Pt(4)
    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.size = Pt(10.5)
    # footer page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(fp, "PAGE")
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("PTM 论文深度分析与专利挖掘研究报告 · 内部研究资料")
    _set_font(r, size=8, color=RGBColor(0x80, 0x80, 0x80))


def _field(p, instr):
    r = p.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r._r.append(b)
    r2 = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = p.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    r4 = p.add_run("1")
    _set_font(r4, size=9)
    r5 = p.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r5._r.append(e)


def add_toc(doc):
    p = doc.add_paragraph()
    _field(p, 'TOC \\o "1-2" \\h \\z \\u')


def render(doc, blocks):
    for b in blocks:
        kind = b[0]
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(b[1], level=int(kind[1]))
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.74)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, b[1])
        elif kind == "pn":  # paragraph without indent
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, b[1])
        elif kind == "note":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            add_inline(p, b[1], size=9, color=RGBColor(0x55, 0x55, 0x55))
        elif kind == "bullets":
            for item in b[1]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, item)
        elif kind == "numbers":
            # manual numbering: every list restarts at 1; step lists ("S1 …") keep their own labels
            for i, item in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.75)
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                label = "" if re.match(r"^S\d", item) else f"{i}. "
                add_inline(p, label + item)
        elif kind == "claims":
            for i, item in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(3)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline(p, f"{i}. {item}", size=10)
        elif kind == "table":
            add_table(doc, b[1], b[2], b[3], size=b[4] if len(b) > 4 else 8.5)
        elif kind == "img":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(b[1]), width=Cm(b[2]))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(cap, b[3], size=9, color=RGBColor(0x40, 0x40, 0x40))
        elif kind == "links":
            for label, text, url in b[1]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.9)
                p.paragraph_format.space_after = Pt(1)
                add_inline(p, f"{label} {text} ", size=9)
                if url:
                    add_hyperlink(p, url, url, size=8.5)
        elif kind == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "toc":
            add_toc(doc)
        else:
            raise ValueError(kind)


def new_doc():
    d = Document()
    setup(d)
    return d
