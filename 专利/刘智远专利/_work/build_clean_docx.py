# -*- coding: utf-8 -*-
"""Rebuild the PTM paper (IEEE TII, 2 columns) into a clean, editable single-column Word file.

- Body text: char-level re-lining (pdf_reflow.py) -> correct word spaces, sub/superscripts,
  de-hyphenation and cross-column paragraph merging.
- Figures and the 8 display equations: 300 dpi crops from the PDF.
- Tables I-III: rebuilt as native, editable Word tables (they are vector outlines in the PDF).
"""
import re
import sys
from collections import Counter
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pdf_reflow import parse_pages, runs_text, to_paragraphs  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")

WORK = Path(__file__).resolve().parent
ROOT = WORK.parent
PDF = next(ROOT.glob("A_Novel_Power_Pylon*.pdf"))
OUT = ROOT / "PTM论文原文_可编辑版.docx"
IMG = WORK / "img"
IMG.mkdir(exist_ok=True)
TEXT_W_IN = 6.5

# display-equation clip windows (page, column, y_top, y_bottom) in PDF points,
# chosen between the neighbouring text lines (see probe_lines.py output)
EQ_WINDOWS = [
    (4, 0, 249.2, 322.2),   # (1) bilateral feature weight
    (5, 1, 705.0, 750.0),   # (2) 3-D restoration
    (6, 0, 548.0, 575.0),   # (3) G = (V, E)
    (6, 0, 716.9, 749.0),   # (4) graph refinement objective
    (6, 1, 328.6, 391.5),   # (5) point-to-segment distance
    (6, 1, 566.3, 605.0),   # (6) 1-DoF diagonal endpoint
    (7, 0, 282.2, 321.0),   # (7) outer-loop stop criterion
    (8, 0, 449.8, 603.5),   # (8) evaluation metrics
]

TABLE1_HEAD = ["Group", "Pylon", "RMSEp", "RMSEn", "NGED^M", "NGED^D", "NGED^A", "PP0.2"]
TABLE1_SUS = """#2 0.135 0.113 0.00 0.02 0.12 0.83
#3 0.129 0.102 0.00 0.02 0.10 0.84
#4 0.092 0.083 0.04 0.00 0.05 0.90
#5 0.173 0.123 0.00 0.05 0.15 0.70
#6 0.130 0.059 0.00 0.00 0.07 0.79
#7 0.144 0.104 0.04 0.01 0.11 0.77
#8 0.089 0.051 0.10 0.01 0.03 0.91
#10 0.106 0.068 0.00 0.02 0.08 0.87
#11 0.117 0.091 0.00 0.03 0.09 0.78
#12 0.139 0.085 0.08 0.06 0.12 0.73
#14 0.112 0.092 0.00 0.01 0.06 0.78
#15 0.131 0.102 0.00 0.00 0.11 0.71
#17 0.099 0.094 0.00 0.02 0.03 0.79
#18 0.108 0.060 0.04 0.03 0.06 0.79
#20 0.140 0.093 0.18 0.09 0.15 0.65
#21 0.158 0.106 0.18 0.08 0.18 0.72
#22 0.123 0.066 0.00 0.02 0.05 0.75
#23 0.133 0.128 0.00 0.04 0.08 0.64
#25 0.115 0.101 0.05 0.00 0.03 0.71
#26 0.167 0.141 0.08 0.06 0.16 0.50
#27 0.106 0.106 0.00 0.00 0.04 0.85
#29 0.140 0.095 0.08 0.05 0.11 0.74
Ave. 0.127 0.094 0.040 0.028 0.090 0.76"""
TABLE1_TEN = """#1 0.212 0.126 0.09 0.06 0.19 0.58
#9 0.157 0.097 0.09 0.04 0.16 0.63
#13 0.180 0.123 0.11 0.09 0.17 0.65
#16 0.134 0.078 0.00 0.02 0.07 0.73
#19 0.156 0.111 0.00 0.01 0.13 0.74
#24 0.164 0.114 0.04 0.03 0.14 0.71
#28 0.203 0.144 0.15 0.11 0.19 0.55
#30 0.117 0.083 0.09 0.08 0.18 0.70
Ave. 0.165 0.110 0.071 0.055 0.154 0.66"""
TABLE1_TOTAL = ["Total Ave.", "", "0.137", "0.098", "0.048", "0.035", "0.107", "0.735"]

TABLE2 = [["Method", "Model characteristic", "Mean RMSEp / m"],
          ["Chen et al.'s", "Abstract-template model", "0.570"],
          ["Zhou et al.'s", "Heuristic contour-based model", "0.501"],
          ["Proposed PTM", "LoD3 truss model", "0.137"]]

TABLE3 = [
    ("Suspension", ["#2", "#3", "#4", "#5", "#6", "#7", "#8", "#10", "#11", "#12", "#14"],
     ["0.65", "0.65", "0.7", "0.74", "0.6", "0.66", "0.55", "0.68", "0.65", "0.73", "0.68"],
     ["0.47", "0.46", "0.51", "0.56", "0.44", "0.49", "0.43", "0.53", "0.5", "0.58", "0.52"]),
    ("Suspension", ["#15", "#17", "#18", "#20", "#21", "#22", "#23", "#25", "#26", "#27", "#29"],
     ["0.75", "0.68", "0.68", "0.6", "0.65", "0.75", "0.7", "0.56", "0.56", "0.51", "0.65"],
     ["0.61", "0.55", "0.54", "0.43", "0.49", "0.59", "0.54", "0.43", "0.41", "0.37", "0.53"]),
    ("Tension", ["#1", "#9", "#13", "#16", "#19", "#24", "#28", "#30"],
     ["0.51", "0.52", "0.55", "0.54", "0.57", "0.53", "0.52", "0.53"],
     ["0.36", "0.36", "0.37", "0.35", "0.4", "0.34", "0.35", "0.35"]),
]


# ----------------------------------------------------------------------------- docx helpers
def add_formatted(p, text, bold=False, size=None):
    """Tiny markup: 'RMSEp' -> RMSE_p, 'NGED^M' -> NGED^M, 'PP0.2' -> PP_0.2."""
    m = re.fullmatch(r"(Mean )?(RMSE)([pn])(.*)", text) or re.fullmatch(r"()(PP)(0\.2)(.*)", text)
    if m:
        parts = [((m.group(1) or "") + m.group(2), None), (m.group(3), "sub"), (m.group(4), None)]
    elif "^" in text:
        a, b = text.split("^", 1)
        parts = [(a, None), (b, "sup")]
    else:
        parts = [(text, None)]
    for t, v in parts:
        if not t:
            continue
        r = p.add_run(t)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if v == "sub":
            r.font.subscript = True
        elif v == "sup":
            r.font.superscript = True


def set_cell(cell, text, bold=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, shade=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    add_formatted(p, text, bold=bold, size=size)
    if shade:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), shade)
        tcPr.append(shd)


def set_table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in ("top", "bottom"):
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "10")
        elif edge == "insideH":
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "2")
            el.set(qn("w:color"), "BFBFBF")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def add_table1(d):
    t = d.add_table(rows=1, cols=len(TABLE1_HEAD))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(TABLE1_HEAD):
        set_cell(t.rows[0].cells[i], h, bold=True, shade="E7EEF6")
    for group, data in (("Suspension", TABLE1_SUS), ("Tension", TABLE1_TEN)):
        for row in data.splitlines():
            vals = row.split()
            cells = t.add_row().cells
            set_cell(cells[0], group if vals[0] in ("#2", "#1") else "")
            for j, v in enumerate(vals):
                set_cell(cells[j + 1], v, bold=(vals[0] == "Ave."))
    cells = t.add_row().cells
    for j, v in enumerate(TABLE1_TOTAL):
        set_cell(cells[j], v, bold=True, shade="F2F2F2")
    set_table_borders(t)


def add_table2(d):
    t = d.add_table(rows=0, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(TABLE2):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            set_cell(cells[j], v, bold=(i == 0), size=9, shade="E7EEF6" if i == 0 else None)
    set_table_borders(t)


def add_table3(d):
    t = d.add_table(rows=0, cols=13)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for group, pylons, tol, dfm in TABLE3:
        for label, vals in (("Pylon", pylons), ("TOL.", tol), ("DEF.", dfm)):
            cells = t.add_row().cells
            head = label == "Pylon"
            set_cell(cells[0], group if head else "", size=7.5)
            set_cell(cells[1], label, bold=True, size=7.5, shade="E7EEF6" if head else None)
            for j, v in enumerate(vals):
                set_cell(cells[j + 2], v, size=7.5, bold=head, shade="E7EEF6" if head else None)
    set_table_borders(t)


def add_runs(p, runs, base_size=None, force_bold=None, color=None):
    for text, (bold, italic, vert) in runs:
        text = re.sub(r"\s{2,}", " ", text)
        r = p.add_run(text)
        r.bold = force_bold if force_bold is not None else bold
        r.italic = italic
        if vert == "sub":
            r.font.subscript = True
        elif vert == "sup":
            r.font.superscript = True
        if base_size:
            r.font.size = Pt(base_size)
        if color:
            r.font.color.rgb = color


def strip_runs(runs):
    if runs:
        runs[0] = (runs[0][0].lstrip(), runs[0][1])
        runs[-1] = (runs[-1][0].rstrip(), runs[-1][1])
    return runs


def build_docx(paras):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 13), ("Heading 2", 11.5)):
        hs = d.styles[name]
        hs.font.name = "Arial"
        hs.font.size = Pt(size)
        hs.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    note = d.add_paragraph()
    r = note.add_run("【说明】本文件由原 PDF 重排生成（单栏可编辑）。正文为逐字形提取并修复词间空格与上下标；"
                     "图与 8 个独立公式为原文 300 dpi 高清截图；表 I–III 为可编辑重建表格（数值逐项核对原文）。"
                     "版权归 IEEE 及原作者所有，仅供个人研究使用。")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    table_map = {"TABLE I": add_table1, "TABLE II": add_table2, "TABLE III": add_table3}
    for p in paras:
        t = p["type"]
        if t in ("figure", "equation"):
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            scale = 1.9 if t == "figure" else 1.15
            para.add_run().add_picture(str(p["path"]), width=Inches(min(TEXT_W_IN, p["w_pt"] / 72 * scale)))
            para.paragraph_format.keep_with_next = t == "figure"
            continue
        runs = strip_runs(p["runs"])
        text = runs_text(runs)
        para = None
        if t == "title":
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(para, runs, base_size=18, force_bold=True)
            for rr in para.runs:
                rr.font.name = "Arial"
        elif t == "authors":
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(para, [(re.sub(r"\s+,", ",", tx), s_) for tx, s_ in runs], base_size=11)
        elif t == "abstract":
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(para, runs, base_size=10)
        elif t in ("h1", "h2"):
            # small-caps glyph gap in the PDF splits this one heading word
            text = {"III. METH OD": "III. METHOD"}.get(text, text)
            d.add_heading(text, level=1 if t == "h1" else 2)
        elif t == "caption":
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(para, runs, base_size=9)
            if para.runs:
                para.runs[0].bold = True
            key = re.match(r"TABLE [IVX]+", text)
            if key and key.group(0) in table_map:
                para.paragraph_format.keep_with_next = True
                table_map[key.group(0)](d)
                d.add_paragraph()
        elif t == "footnote":
            para = d.add_paragraph()
            add_runs(para, runs, base_size=8.5, color=RGBColor(0x40, 0x40, 0x40))
        elif t == "ref":
            para = d.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.8)
            para.paragraph_format.first_line_indent = Cm(-0.8)
            para.paragraph_format.space_after = Pt(2)
            add_runs(para, runs, base_size=9)
        else:
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.first_line_indent = Cm(0.6)
            add_runs(para, runs)
    d.save(OUT)


def main():
    doc = pymupdf.open(PDF)
    items = parse_pages(doc, EQ_WINDOWS, IMG)
    paras = to_paragraphs(items)
    with open(WORK / "clean_dump.txt", "w", encoding="utf-8") as fh:
        for p in paras:
            if p["type"] in ("figure", "equation"):
                fh.write(f"<<{p['type']} {p['path'].name}>>\n\n")
            else:
                fh.write(f"[{p['type']}] {runs_text(p['runs']).strip()}\n\n")
    build_docx(paras)
    print("items:", dict(Counter(p["type"] for p in paras)))
    print("saved:", OUT.name, round(OUT.stat().st_size / 1e6, 2), "MB")


if __name__ == "__main__":
    main()
